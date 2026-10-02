"""
document_service.py — file I/O and document processing for StudySense.

Phases covered:
  Phase 1 — PDF upload and ingestion
  Phase 3 — TXT and DOCX support; user-scoped ChromaDB collections

Responsibilities:
  - Validate uploaded files (extension + MIME type)
  - Persist uploads to disk with UUID-based filenames
  - Create/update Document rows in SQLite
  - Dispatch to the correct document loader (PDF / TXT / DOCX)
  - Chunk, embed, and persist in ChromaDB

Heavy LangChain / ChromaDB imports are deferred inside ``process_document``
so the module starts quickly.
"""

import uuid
from pathlib import Path
from typing import Optional

from fastapi import HTTPException, UploadFile
from sqlalchemy.orm import Session

from backend.config.settings import settings
from backend.database.models import Document


# ── Constants ─────────────────────────────────────────────────────────────────

_ALLOWED_EXTENSIONS: dict[str, str] = {
    ".pdf": "pdf",
    ".txt": "txt",
    ".docx": "docx",
}

# Content-type check is intentionally lenient: many HTTP clients send
# application/octet-stream regardless of the real file type.
_ALLOWED_CONTENT_TYPES: set[str] = {
    "application/pdf",
    "text/plain",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "application/octet-stream",
    "binary/octet-stream",
}


# ── Validation ────────────────────────────────────────────────────────────────

def validate_file(file: UploadFile) -> str:
    """
    Validate the uploaded file extension and MIME type.

    Returns the file_type string ("pdf", "txt", "docx").
    Raises HTTPException(400) on invalid input.
    """
    original_name = file.filename or ""
    suffix = Path(original_name).suffix.lower()

    if suffix not in _ALLOWED_EXTENSIONS:
        raise HTTPException(
            status_code=400,
            detail=(
                f"Unsupported file type '{suffix or '(none)'}'. "
                f"Accepted: {', '.join(_ALLOWED_EXTENSIONS)}"
            ),
        )

    content_type = file.content_type or ""
    if content_type and content_type not in _ALLOWED_CONTENT_TYPES:
        raise HTTPException(
            status_code=400,
            detail=f"Unexpected content-type '{content_type}'.",
        )

    return _ALLOWED_EXTENSIONS[suffix]


# ── File I/O ──────────────────────────────────────────────────────────────────

def save_upload_file(file: UploadFile) -> tuple[str, str, str]:
    """
    Persist the uploaded file to the configured upload directory.

    A UUID-based filename is generated to prevent collisions and path-
    traversal vulnerabilities from untrusted original filenames.

    Returns:
        (original_filename, stored_filename, stored_path)
    """
    upload_dir = Path(settings.upload_dir)
    upload_dir.mkdir(parents=True, exist_ok=True)

    original_filename = file.filename or "upload"
    suffix = Path(original_filename).suffix.lower()
    stored_filename = f"{uuid.uuid4().hex}{suffix}"
    stored_path = str(upload_dir / stored_filename)

    content = file.file.read()
    if not content:
        raise HTTPException(status_code=400, detail="Uploaded file is empty.")

    with open(stored_path, "wb") as fh:
        fh.write(content)

    return original_filename, stored_filename, stored_path


# ── Database helpers ──────────────────────────────────────────────────────────

def create_document_record(
    db: Session,
    *,
    filename: str,
    stored_filename: str,
    stored_path: str,
    file_type: str,
    user_id: Optional[int] = None,
) -> Document:
    """Insert a new Document row (is_processed=False) and return it."""
    doc = Document(
        filename=filename,
        stored_filename=stored_filename,
        stored_path=stored_path,
        file_type=file_type,
        is_processed=False,
        chunk_count=0,
        user_id=user_id,
    )
    db.add(doc)
    db.commit()
    db.refresh(doc)
    return doc


# ── Document loaders ──────────────────────────────────────────────────────────

def _load_pdf(file_path: str) -> list:
    """Load a PDF with PyMuPDF via LangChain."""
    from langchain_community.document_loaders import PyMuPDFLoader

    loader = PyMuPDFLoader(file_path)
    docs = loader.load()
    if not docs:
        raise ValueError(
            "PyMuPDFLoader returned no pages — "
            "the PDF may be empty or purely image-based."
        )
    return docs


def _load_txt(file_path: str, filename: str) -> list:
    """
    Read a UTF-8 text file and wrap it in a single LangChain Document.
    Falls back to replacing undecodable bytes rather than crashing.
    """
    from langchain_core.documents import Document as LCDocument

    with open(file_path, "r", encoding="utf-8", errors="replace") as fh:
        text = fh.read()

    if not text.strip():
        raise ValueError("TXT file is empty or contains no readable text.")

    return [LCDocument(page_content=text, metadata={"source": filename})]


def _load_docx(file_path: str, filename: str) -> list:
    """
    Extract paragraph text from a DOCX file using python-docx.
    Returns a single LangChain Document containing all non-empty paragraphs.
    """
    from docx import Document as DocxDoc
    from langchain_core.documents import Document as LCDocument

    word_doc = DocxDoc(file_path)
    paragraphs = [p.text.strip() for p in word_doc.paragraphs if p.text.strip()]

    if not paragraphs:
        raise ValueError("DOCX file contains no extractable text paragraphs.")

    text = "\n\n".join(paragraphs)
    return [LCDocument(page_content=text, metadata={"source": filename})]


def _collection_name(user_id: Optional[int], doc_id: int) -> str:
    """
    Generate a ChromaDB collection name.

    Format:
      New docs:    studysense_user_{user_id}_doc_{doc_id}
      Legacy docs: studysense_doc_{doc_id}
    """
    if user_id is not None:
        return f"studysense_user_{user_id}_doc_{doc_id}"
    return f"studysense_doc_{doc_id}"


# ── Processing pipeline ───────────────────────────────────────────────────────

def process_document(db: Session, doc: Document) -> Document:
    """
    Run the full ingestion pipeline:

    1. Load the file (PDF / TXT / DOCX)
    2. Split into chunks (RecursiveCharacterTextSplitter)
    3. Attach doc_id, user_id, filename metadata to every chunk
    4. Embed with OllamaEmbeddings
    5. Persist in a dedicated ChromaDB collection

    Updates ``collection_name``, ``chunk_count``, ``is_processed=True`` on
    success.  Raises HTTPException(503) on failure without marking as processed.
    """
    from langchain_text_splitters import RecursiveCharacterTextSplitter
    from langchain_ollama import OllamaEmbeddings
    from langchain_chroma import Chroma

    collection_name = _collection_name(doc.user_id, doc.id)

    try:
        # ── 1. Load ───────────────────────────────────────────────────────
        if doc.file_type == "pdf":
            raw_docs = _load_pdf(doc.stored_path)
        elif doc.file_type == "txt":
            raw_docs = _load_txt(doc.stored_path, doc.filename)
        elif doc.file_type == "docx":
            raw_docs = _load_docx(doc.stored_path, doc.filename)
        else:
            raise ValueError(f"Unsupported file_type '{doc.file_type}'.")

        # ── 2. Split ──────────────────────────────────────────────────────
        splitter = RecursiveCharacterTextSplitter(
            chunk_size=settings.chunk_size,
            chunk_overlap=settings.chunk_overlap,
            add_start_index=True,
        )
        chunks = splitter.split_documents(raw_docs)

        if not chunks:
            raise ValueError("Text splitter produced zero chunks.")

        # ── 3. Enrich chunk metadata ──────────────────────────────────────
        for chunk in chunks:
            chunk.metadata["doc_id"] = doc.id
            chunk.metadata["user_id"] = doc.user_id
            chunk.metadata["filename"] = doc.filename
            chunk.metadata["file_type"] = doc.file_type

        # ── 4. Embed ──────────────────────────────────────────────────────
        embeddings = OllamaEmbeddings(
            model=settings.ollama_embedding_model,
            base_url=settings.ollama_base_url,
        )

        # ── 5. Persist ────────────────────────────────────────────────────
        Chroma.from_documents(
            documents=chunks,
            embedding=embeddings,
            collection_name=collection_name,
            persist_directory=settings.vector_db_dir,
        )

    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(
            status_code=503,
            detail=(
                f"Document processing failed: {exc}  "
                "Ensure Ollama is running and the configured embedding model is available."
            ),
        ) from exc

    # ── 6. Mark as processed ──────────────────────────────────────────────
    doc.collection_name = collection_name
    doc.chunk_count = len(chunks)
    doc.is_processed = True
    db.commit()
    db.refresh(doc)
    return doc
